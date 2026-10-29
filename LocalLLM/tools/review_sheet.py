"""Alert review workbook for a LEONI expert, in LEONI's own form IT 3535-11 Anexo 4 "Requisitos clientes - análise da
conformidade" (template complyai/templates/IT-3535-11-A4_template.xlsx, made by tools/make_a4_template.py). The sheet
"Verification" lists a sample of the AI's alerts (every conflict + a random, per-document share of the gaps) as
requirement rows under one heading row per customer document, with the AI's proposal in the form's own fields
(comments, action plan, action status, the IT / procedure that reflects the requirement, STATUS); the expert checks
and corrects them like any conformity analysis. Grey columns keep the AI's original proposal; two calculated columns
turn the expert's final STATUS and IT into "AI alert confirmed?" / "AI's IT confirmed?", counted on RESUMO next to
LEONI's own summary and action plan. Extra sheets: Read me, All alerts, Sources. No LLM: reads the database only.

    python tools/review_sheet.py [--gaps 26] [--seed 2026] [--out data/exports/review/...xlsx]
"""
import argparse
import datetime
import json
import math
import random
import re
import sqlite3
import sys
from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "complyai" / "templates" / "IT-3535-11-A4_template.xlsx"
sys.stdout.reconfigure(encoding="utf-8")
ap = argparse.ArgumentParser()
ap.add_argument("--db", default=str(ROOT / "data" / "complyai.db"))
ap.add_argument("--gaps", type=int, default=26, help="gaps in the sample (all conflicts are always included)")
ap.add_argument("--seed", type=int, default=2026)
ap.add_argument("--out", default=str(ROOT / "data" / "exports" / "review"
                                    / f"ComplyAI_alert_review_{datetime.date.today():%Y-%m-%d}.xlsx"))
a = ap.parse_args()

# LEONI's STATUS values (sheet "Dop List") that confirm an alert: the requirement is not (fully) met
CONFIRMS = ("Do not comply", "Action required")
AI_STATUS = {"Conflict": "Do not comply", "Gap": "Action required"}

db = sqlite3.connect(f"file:{a.db}?mode=ro", uri=True)
db.row_factory = sqlite3.Row
docs = {d["id"]: d for d in db.execute("SELECT * FROM documents WHERE kind='CSR' AND status='active'")}
its = {d["code"]: d for d in db.execute("SELECT * FROM documents WHERE kind='IT' AND status='active'")}
labels = {(r["doc_id"], r["section"]): r["label"] for r in db.execute("SELECT * FROM section_labels WHERE language='English'")}
rows = []
for f in db.execute("""SELECT f.*, r.page, r.section, r.statement, r.quote, l.evidence, l.explanation
                       FROM findings f JOIN requirements r ON r.id=f.req_id
                       LEFT JOIN req_it_links l ON l.req_id=f.req_id AND l.it_code=f.it_code AND l.source='text'
                       WHERE f.type IN ('CONFLICT','GAP') AND f.doc_id IN (%s)
                       ORDER BY f.doc_id, r.page, f.id""" % ",".join(map(str, docs))):
    d, it = docs[f["doc_id"]], its.get(f["it_code"])
    rows.append({"id": f["id"], "type": "Conflict" if f["type"] == "CONFLICT" else "Gap", "doc_id": f["doc_id"],
                 "client": d["client"], "document": f"{d['title']} ({d['code']}, {d['version'] or '-'})",
                 "page": f["page"], "section": labels.get((f["doc_id"], f["section"]), f["section"] or ""),
                 "statement": f["statement"], "quote": f["quote"], "it_code": f["it_code"],
                 "it": f"{f['it_code']} {it['title'] if it else ''}".strip(), "evidence": f["evidence"] or "",
                 "explanation": f["explanation"] or "", "action": f["action"] or ""})

# sample: every conflict + gaps per document in proportion to its gaps (at least 2), reproducible
rng = random.Random(a.seed)
gaps_by_doc = {k: [r for r in rows if r["doc_id"] == k and r["type"] == "Gap"] for k in docs}
n_gaps = sum(map(len, gaps_by_doc.values()))
quota = {k: min(len(g), max(2, round(a.gaps * len(g) / n_gaps))) for k, g in gaps_by_doc.items()}
while sum(quota.values()) > a.gaps:  # rounding + minimums: trim the largest documents
    quota[max(quota, key=lambda k: (quota[k], len(gaps_by_doc[k])))] -= 1
picked = {r["id"] for r in rows if r["type"] == "Conflict"}
for k, g in gaps_by_doc.items():
    picked |= {r["id"] for r in rng.sample(g, quota[k])}
sample = [r for r in rows if r["id"] in picked]
n_conf = sum(r["type"] == "Conflict" for r in rows)
n_conf_sample = sum(r["type"] == "Conflict" for r in sample)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
AI_HEAD, AI_CELL = PatternFill("solid", fgColor="BFBFBF"), PatternFill("solid", fgColor="F2F2F2")
CALC_HEAD, CALC_CELL = PatternFill("solid", fgColor="BDD7EE"), PatternFill("solid", fgColor="DDEBF7")
CONF_FILL, GAP_FILL = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFE699")


def style_of(c):
    return {k: copy(getattr(c, k)) for k in ("font", "border", "fill", "alignment", "number_format", "protection")}


def apply(c, st):
    for k, v in st.items():
        setattr(c, k, copy(v))


def lines(text, width, size):
    """Wrapped lines of `text` in a column `width` characters wide (Calibri 11 units) at font `size`."""
    per_line = max(int(width * 11 / size * 1.05), 8)
    return sum(max(1, math.ceil(len(p) / per_line)) for p in str(text or "").split("\n"))


wb = load_workbook(TEMPLATE)
ws = wb["Verification"]
head_style = {c: style_of(ws.cell(5, c)) for c in range(1, 19)}  # LEONI's section heading row
row_style = {c: style_of(ws.cell(6, c)) for c in range(1, 19)}   # LEONI's requirement row
clients = " / ".join(dict.fromkeys(d["client"].upper() for _, d in sorted(docs.items())))
ws["F1"] = clients
ws["F2"] = f"COMPLY-AI alert review: {len(sample)} alerts from {len(docs)} customer documents (sheet 'Sources')"
ws["F3"] = f"Prepared {datetime.date.today():%d.%m.%Y}"
ws.merge_cells("J1:Q3")
note = ws["J1"]
note.value = ("FILLED IN BY COMPLY-AI - PLEASE CHECK AND CORRECT. Each row is an alert raised by the AI: it filled in "
              "comments (J), action plan (K), the IT / procedure (O / Q) and STATUS (R). Correct whatever is wrong, as in "
              "a normal conformity analysis; if the requirement is met, set STATUS to Complies (or Agreed / Not "
              "applicable) and clear the action plan. Grey columns S-X keep the AI's original proposal, blue Y-Z compare "
              "it with your answer (totals on RESUMO). Details: sheet 'Read me'.")
note.font, note.alignment = Font(name="Calibri", size=14, italic=True, color="404040"), WRAP

AI_COLS = [("S", "COMPLY-AI\nALERT", 16), ("T", "AI STATUS\n(as proposed)", 22), ("U", "AI IT\n(as proposed)", 36),
           ("V", "REQUIREMENT AS READ\nBY THE AI (EN)", 60), ("W", "ALERT\nID", 10), ("X", "CUSTOMER\nDOCUMENT", 34)]
CALC_COLS = [("Y", "AI ALERT\nCONFIRMED?\n(from R)", 20), ("Z", "AI'S IT\nCONFIRMED?\n(from O / Q)", 20)]
CENTRAL = ("VA", "PI")  # central LEONI procedures go to "Reflected LEONI IMS" (O), local ITs to "Reflected WPTGU" (Q)


def it_label(code):
    """'WI 3157-11' -> 'IT 3157-11 Formação / Integração de colaboradores' (LEONI writes local WIs as IT)."""
    prefix, rest = code.split(" ", 1)
    title = its[code]["title"] if code in its else ""
    return f"{'IT' if prefix == 'WI' else prefix} {rest} {title}".strip()
header = style_of(ws["R4"])
for col, title, width in AI_COLS + CALC_COLS:
    c = ws[f"{col}4"]
    c.value = title
    apply(c, header)
    c.font = Font(name="Calibri", size=16, bold=True)
    c.fill = CALC_HEAD if (col, title, width) in CALC_COLS else AI_HEAD
    ws.column_dimensions[col].width = width
ws.column_dimensions["D"].width = 26  # LEONI writes the section letter here; COMPLY-AI writes the section heading

r = 5
for doc_id, d in sorted(docs.items()):
    alerts = [x for x in sample if x["doc_id"] == doc_id]
    if not alerts:
        continue
    for c in range(1, 19):
        apply(ws.cell(r, c), head_style[c])
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    ws.cell(r, 5, f"{d['client']} - {d['title']} ({d['code']}, {d['version'] or '-'}): {len(alerts)} alerts")
    ws.row_dimensions[r].height = 36
    ws.cell(r, 1, f"=IF(K{r}=0,0,1)")
    ws.cell(r, 2, f"=A{r}" if r == 5 else f"=B{r - 1}+A{r}")
    r += 1
    for x in alerts:
        for c in range(1, 19):
            apply(ws.cell(r, c), row_style[c])
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        ws.cell(r, 1, f"=IF(K{r}=0,0,1)")
        ws.cell(r, 2, f"=B{r - 1}+A{r}")
        # COMPLY-AI's proposal goes into the form's own fields; the reviewer corrects it like any analysis
        clause = re.search(r"\((\d+(?:\.\d+)+)\)", x["section"] or "")       # 'M. ... (8.2.1)' -> ISO 9001 clause
        evidence = f'\nIT text ({x["it_code"]}): "{x["evidence"]}"' if x["evidence"] else ""
        fields = {3: x["page"], 4: x["section"], 5: x["quote"] or x["statement"], 8: clause.group(1) if clause else None,
                  10: f"COMPLY-AI ({x['type'].lower()}): {x['explanation']}{evidence}", 11: x["action"] or None,
                  14: "TO BE DETERMINED" if x["action"] else None,
                  15 if x["it_code"].split()[0] in CENTRAL else 17: it_label(x["it_code"]), 18: AI_STATUS[x["type"]]}
        for col, v in fields.items():
            ws.cell(r, col, v)
        ws.cell(r, 4).alignment = WRAP
        values = [x["type"], AI_STATUS[x["type"]], x["it"], x["statement"], x["id"],
                  f"{x['client']} {docs[x['doc_id']]['code']}"]
        for (col, _, _), v in zip(AI_COLS, values):
            c = ws[f"{col}{r}"]
            c.value, c.font, c.alignment, c.border, c.fill = v, Font(name="Calibri", size=16), WRAP, BORDER, AI_CELL
        ws[f"S{r}"].fill = CONF_FILL if x["type"] == "Conflict" else GAP_FILL
        num = re.search(r"\d\s?\d{3}", x["it_code"]).group().replace(" ", "")  # 'WI 3326-11 A3' -> '3326'
        ws[f"Y{r}"] = f'=IF(R{r}="","",IF(OR(R{r}="{CONFIRMS[0]}",R{r}="{CONFIRMS[1]}"),"Yes","No"))'
        ws[f"Z{r}"] = f'=IF(AND(O{r}="",Q{r}=""),"",IF(ISNUMBER(SEARCH("{num}",O{r}&" "&Q{r})),"Yes","No"))'
        for col in ("Y", "Z"):
            c = ws[f"{col}{r}"]
            c.font, c.alignment, c.border, c.fill = Font(name="Calibri", size=16, bold=True), \
                Alignment(horizontal="center", vertical="center"), BORDER, CALC_CELL
        widths = {5: 170, 4: 26, 10: 51, 11: 89, 15: 42, 17: 40} | {i: w for i, (_, _, w) in enumerate(AI_COLS, start=19)}
        texts = {i: ws.cell(r, i).value for i in (4, 5, 10, 11, 15, 17)} | {i: v for i, v in enumerate(values, start=19)}
        ws.row_dimensions[r].height = min(409, max(60, max(lines(texts[i], widths[i], 16) for i in texts) * 21 + 10))
        r += 1
last = r - 1

for col, src in (("R", "'Dop List'!$D$3:$D$9"), ("N", "'Dop List'!$D$13:$D$16")):  # LEONI's own drop-down lists
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    dv.add(f"{col}5:{col}{last}")
    ws.add_data_validation(dv)
ws.freeze_panes = "E5"
ws.auto_filter.ref = f"C4:Z{last}"
ws.print_area = f"C1:Z{last}"

# ---- RESUMO: LEONI's summary (from STATUS / action plan) + the COMPLY-AI review counts ----------------------
rs = wb["RESUMO"]
label, value, pct = style_of(rs["B8"]), style_of(rs["D8"]), style_of(rs["D6"])
for col, width in (("L", 46), ("M", 12), ("N", 12)):
    rs.column_dimensions[col].width = width
rs["L6"] = "COMPLY-AI ALERT REVIEW"
apply(rs["L6"], style_of(rs["B23"]))
ab, ac, s = f"Verification!$Y$5:$Y${last}", f"Verification!$Z$5:$Z${last}", f"Verification!$S$5:$S${last}"
counts = [("AI alerts in this review", f"=COUNTA(Verification!$W$5:$W${last})", None),
          ("Alerts with a STATUS (R)", f'=COUNTIF({ab},"Yes")+COUNTIF({ab},"No")', None),
          ("AI alert confirmed (STATUS still Do not comply / Action required)", f'=COUNTIF({ab},"Yes")', "M10"),
          ("  of the conflicts", f'=COUNTIFS({s},"Conflict",{ab},"Yes")',
           f'COUNTIFS({s},"Conflict",{ab},"Yes")+COUNTIFS({s},"Conflict",{ab},"No")'),
          ("  of the gaps", f'=COUNTIFS({s},"Gap",{ab},"Yes")',
           f'COUNTIFS({s},"Gap",{ab},"Yes")+COUNTIFS({s},"Gap",{ab},"No")'),
          ("Alerts with an IT / procedure named in O or Q", f'=COUNTIF({ac},"Yes")+COUNTIF({ac},"No")', None),
          ("AI compared the right IT (its IT number still in O or Q)", f'=COUNTIF({ac},"Yes")', "M18")]
for i, (text, formula, share_of) in enumerate(counts):
    row = 8 + 2 * i
    rs[f"L{row}"], rs[f"M{row}"] = text, formula
    apply(rs[f"L{row}"], label)
    apply(rs[f"M{row}"], value)
    if share_of:
        rs[f"N{row}"] = f"=IFERROR(M{row}/{share_of},\"-\")" if share_of.startswith("M") else \
            f"=IFERROR(M{row}/({share_of}),\"-\")"
        apply(rs[f"N{row}"], pct)

# LEONI's action plan (PLANO DE AÇÕES) has 25 lines; the AI proposes an action on most alerts, so add lines with the
# same lookups (action k = k-th row of Verification with an action plan in K)
n_actions = sum(1 for x in sample if x["action"])
proto = {c: style_of(rs.cell(49, c)) for c in range(1, 11)}
for k in range(26, n_actions + 1):
    row = 24 + k
    for c in range(1, 11):
        apply(rs.cell(row, c), proto[c])
    rs.row_dimensions[row].height = rs.row_dimensions[49].height
    rs.cell(row, 1, k)
    for col, letter, idx in (("B", "D", 3), ("D", "K", 10), ("F", "L", 11), ("H", "M", 12), ("I", "N", 13)):
        rng_ = f"Verification!$B$5:${letter}${last}"
        rs[f"{col}{row}"] = f'=IF(ISERROR(VLOOKUP(A{row},{rng_},{idx},FALSE)),"",(VLOOKUP(A{row},{rng_},{idx},FALSE)))'
    for a_, b_ in (("B", "C"), ("D", "E"), ("I", "J")):
        rs.merge_cells(f"{a_}{row}:{b_}{row}")
rs["H10"] = f"=MAX(Verification!$B$5:$B${last})"  # TOTAL OF ACTIONS (COUNTA over the lookup lines counted all 25)
rs.print_area = f"A1:N{24 + max(25, n_actions)}"

# ---- Read me ------------------------------------------------------------------------------------------------
me = wb.create_sheet("Read me")
me.column_dimensions["A"].width = 130
FONT, BOLD = Font(name="Calibri", size=11), Font(name="Calibri", size=11, bold=True)
text = [
    ("LEONI COMPLY-AI - review of AI alerts (form IT 3535-11 Anexo 4)", Font(name="Calibri", size=16, bold=True)),
    (f"Prepared {datetime.date.today():%d.%m.%Y}. Internal: contains text of customer documents and of LEONI work "
     "instructions.", Font(name="Calibri", size=10, italic=True, color="7F7F7F")),
    ("", None),
    ("What this is", BOLD),
    ("COMPLY-AI is a prototype that reads customer requirement documents (CSR / standards), extracts each requirement "
     "and compares it with LEONI Portugal's work instructions (ITs). When an IT contradicts a requirement it raises a "
     "CONFLICT; when an IT deals with the topic but does not fully cover the requirement it raises a GAP. A LEONI expert's "
     "assessment of a sample of these alerts tells us how often they are right.", FONT),
    (f"Analysed: {len(docs)} customer documents ({sum(1 for _ in db.execute('SELECT 1 FROM requirements WHERE status=?', ('active',)))} "
     f"requirements) against {len(its)} LEONI IT documents (about 15 ITs with their annexes, of ~300; list on 'Sources'). "
     f"Result: {len(rows)} alerts ({n_conf} conflicts, {len(rows) - n_conf} gaps), all on 'All alerts'. Because the AI "
     "knows so few ITs, it often compares a requirement with the closest IT it has, which may not be the one that "
     "governs it; this review should show how often.", FONT),
    ("", None),
    ("What to do (sheet 'Verification', the usual conformity-analysis form, filled in by COMPLY-AI)", BOLD),
    (f"{len(sample)} alerts: every conflict ({n_conf_sample}; there are few and they are the most serious) plus "
     f"{len(sample) - n_conf_sample} of the {len(rows) - n_conf} gaps, picked at random and spread over the documents in "
     "proportion to their gaps (at least 2 per document), under one heading row per customer document. About 1-2 minutes "
     "each. COMPLY-AI filled in each row as a proposal:", FONT),
    ("  C-E  page, section and the customer's own wording; H  the ISO 9001 clause where the customer's heading gives one.",
     FONT),
    ("  J  Comments: the AI's analysis (starting 'COMPLY-AI (gap)' or 'COMPLY-AI (conflict)'; for a gap usually what "
     "the IT does not cover) and the IT text it relied on. K  Action plan: the action "
     "it recommends. N  ACTION STATUS: TO BE DETERMINED.", FONT),
    ("  Q / O  Reflected WPTGU / LEONI IMS: the local IT (Q) or central procedure (O) it compared the requirement with.",
     FONT),
    ("  R  STATUS: 'Do not comply' for a CONFLICT (the IT contradicts the requirement), 'Action required' for a GAP (the "
     "IT does not fully cover it).", FONT),
    ("Please check every row and correct what is wrong, as in any analysis: set STATUS (R) to your own assessment "
     "(Complies, Agreed, Do not comply, Not assessed, Action required, Acknowledged, Not applicable); if the requirement "
     "is met, also clear the action plan (K) and ACTION STATUS (N) so it leaves RESUMO's action plan; replace the IT in "
     "Q / O if another IT or procedure governs the requirement; fill RESPONSIBLE (L) and DEADLINE (M) for the actions "
     "you keep; edit the comments (J) and actions (K) freely. '[not verbatim]' in the IT text marks wording the AI wrote "
     "itself because no matching IT sentence was found; '[...]' joins two sentences of the same IT.", FONT),
    ("Grey columns S-X keep the AI's original proposal (alert type, status, IT, how it read the requirement in English, "
     "alert ID, document), so your corrections can be compared with it. Blue columns Y-Z (calculated, do not fill in): "
     "Y = the alert is confirmed when your STATUS is still 'Do not comply' or 'Action required'; Z = the AI compared the "
     "right IT when its IT number is still in Q or O. RESUMO counts both (columns L-N), next to LEONI's usual summary "
     "and action plan.", FONT),
    ("", None),
    ("Example of a corrected row (illustration only, not one of the alerts)", BOLD),
]
for i, (t, font) in enumerate(text, start=1):
    c = me.cell(i, 1, t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font
r0 = len(text) + 1
example = [("AI proposal", "Q = IT 3326-11 A3, R = Action required, K = 'Add the Ppk >= 1.67 criterion to the IT'"),
           ("R  STATUS (corrected)", "Complies"),
           ("Q  Reflected WPTGU (corrected)", "IT 3519-11"),
           ("J  Comments (added)", "Ppk is handled in the control plan per IT 3519-11, not in the sample submission IT."),
           ("K / N  (cleared)", "no action needed"),
           ("Y / Z (calculated)", "Y = No (alert not confirmed), Z = No (the AI compared another IT)")]
for i, (k, v) in enumerate(example):
    c = me.cell(r0 + i, 1, f"{k}:  {v}")
    c.font, c.fill = FONT, PatternFill("solid", fgColor="FFF2CC")
me.cell(r0 + len(example) + 1, 1, "Please return the file to the sender when done.").font = BOLD


def table(sheet, headers, widths):
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = sheet.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = Font(name="Calibri", size=11, bold=True, color="FFFFFF"), \
            PatternFill("solid", fgColor="1F3864"), WRAP, BORDER
        sheet.column_dimensions[get_column_letter(j)].width = w
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def put(sheet, row, values):
    for j, v in enumerate(values, start=1):
        c = sheet.cell(row, j, v)
        c.font, c.alignment, c.border = FONT, WRAP, BORDER


# ---- All alerts / Sources (reference) -------------------------------------------------------------------------
al = wb.create_sheet("All alerts")
table(al, ["No.", "Alert ID", "Type", "Customer", "Customer document", "Page", "Section",
           "Customer requirement (AI summary, English)", "LEONI IT compared", "IT text the AI relied on",
           "AI analysis", "In review sample"], [5, 7, 9, 11, 22, 6, 22, 50, 26, 42, 50, 9])
for i, x in enumerate(rows, start=1):
    put(al, i + 1, [i, x["id"], x["type"], x["client"], x["document"], x["page"], x["section"], x["statement"],
                    x["it"], x["evidence"], x["explanation"], "Yes" if x["id"] in picked else ""])
    al.cell(i + 1, 3).fill = CONF_FILL if x["type"] == "Conflict" else GAP_FILL

so = wb.create_sheet("Sources")
table(so, ["Customer document", "Customer", "Code", "Version", "Requirements", "Conflicts", "Gaps"],
      [60, 14, 26, 16, 13, 10, 10])
per_doc = Counter((x["doc_id"], x["type"]) for x in rows)
for i, (k, d) in enumerate(sorted(docs.items()), start=2):
    nreq = db.execute("SELECT COUNT(*) FROM requirements WHERE doc_id=? AND status='active'", (k,)).fetchone()[0]
    put(so, i, [d["title"], d["client"], d["code"], d["version"], nreq, per_doc[(k, "Conflict")], per_doc[(k, "Gap")]])
r0 = len(docs) + 4
for j, h in enumerate(["LEONI IT document", "Title", "Version", "Rules read by the AI", "Alerts against it"], start=1):
    c = so.cell(r0, j, h)
    c.font, c.fill, c.border = Font(name="Calibri", size=11, bold=True, color="FFFFFF"), \
        PatternFill("solid", fgColor="1F3864"), BORDER
by_it = Counter(x["it_code"] for x in rows)
for i, (code, d) in enumerate(sorted(its.items()), start=r0 + 1):
    nrules = db.execute("SELECT COUNT(*) FROM it_rules WHERE doc_id=?", (d["id"],)).fetchone()[0]
    put(so, i, [code, d["title"], d["version"], nrules, by_it.get(code, 0)])

wb.move_sheet("Dop List", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Dop List"))  # LEONI keeps it last
for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = False
wb.active = wb.sheetnames.index("Verification")
ws.sheet_view.tabSelected = True
wb.calculation.fullCalcOnLoad = True
wb.properties.creator = wb.properties.lastModifiedBy = "LEONI COMPLY-AI"
wb.properties.title = "COMPLY-AI alert review (IT 3535-11 Anexo 4)"
Path(a.out).parent.mkdir(parents=True, exist_ok=True)
wb.save(a.out)
print(f"{a.out}: {len(sample)} alerts to review ({n_conf_sample} conflicts, {len(sample) - n_conf_sample} gaps; "
      f"gaps per document {dict(quota)}), {len(rows)} alerts listed")
print(json.dumps({"sample_ids": sorted(picked)}))
