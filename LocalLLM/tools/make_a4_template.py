"""Build complyai/templates/IT-3535-11-A4_template.xlsx from LEONI's own form IT 3535-11 Anexo 4 "Requisitos clientes -
análise da conformidade" (filled for the Cummins handbook; the copy in the repo root). Only the layout is kept: sheets
RESUMO (summary + action plan formulas), Verification (header rows 1-4, one heading row and one requirement row as
style prototypes, columns C-R) and Dop List (drop-down values); every Cummins row, answer and picture is removed, and the
LEONI logo stays. LEONI's analysis is never used as a source of answers.

    python tools/make_a4_template.py ["IT-3535-11-A4_Requisitos clientes – análise da conformidade.xlsx"]
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "IT-3535-11-A4_Requisitos clientes – análise da conformidade.xlsx"
OUT = ROOT / "complyai" / "templates" / "IT-3535-11-A4_template.xlsx"
LAST_COL = 18  # R = STATUS; S onwards (scope, LEONI processes, topic ticks) is not used by COMPLY-AI

wb = load_workbook(SRC)
for name in ("VerificationV2025", "Historial", "TABEL P17", "TABLE 1", "rev log"):
    del wb[name]  # revision-change analysis, history and pasted customer tables

for ws in (wb["RESUMO"], wb["Verification"]):
    ws._images = [img for img in ws._images if img.anchor._from.row == 0]  # LEONI logo only, no customer pictures

ws = wb["Verification"]
for rng in list(ws.merged_cells.ranges):
    if rng.min_row >= 5 or rng.max_col > LAST_COL:
        ws.unmerge_cells(str(rng))
ws.delete_rows(7, ws.max_row)                       # rows 5 (section heading) and 6 (requirement) stay as prototypes
ws.delete_cols(LAST_COL + 1, ws.max_column)
for row in ws.iter_rows(min_row=1, max_row=6):
    for c in row:
        if not isinstance(c, MergedCell) and (c.row >= 5 or c.column == 6 or c.column > LAST_COL):
            c.value = None                          # customer name, document, revision and all answers
ws.auto_filter.ref = None
ws.print_area = f"C1:R6"
wb.properties.creator = wb.properties.lastModifiedBy = "LEONI COMPLY-AI"  # not the person who filled the original
wb.properties.title = "IT 3535-11 Anexo 4 template"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"{OUT}: sheets {wb.sheetnames}")
