"""Compliance matrix export (Excel).

Sheet "Matrix" reproduces LEONI's own form IT 3535-11 Anexo 5 "Requisitos Específicos Clientes - Matriz de
correspondência" (template in complyai/templates/, made from LEONI's Cummins matrix): every section of the
customer document with its page, the LEONI process, the central LEONI procedures loaded in the system (VA / PI
codes, PROC. WSD) and the local ITs (PROC. WPTGU). Most central procedures are not in the system: PROC. WSD is
completed by hand.

The extra sheets carry what the form has no place for: "Status by section" (verdicts per section),
"Requirements" (one row per requirement with evidence and action) and "Alerts".
"""
import copy
import datetime
import json
import math
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .engine import section_labels
from .parser import MATRIX_LEVEL
from .store import Store

TEMPLATE = Path(__file__).parent / "templates" / "IT-3535-11-A5_template.xlsx"
FIRST_ROW = 9          # first data row of the template (row 8 = column headers)
COL = {"section": 2, "page": 3, "process": 5, "wsd": 7, "wptgu": 10}   # B, C, E, G, J as in LEONI's form

STATUS = {  # verdict -> (label, fill colour, rank)
    "conflict": ("Conflict", "F8CBAD", 4),
    "partial": ("Partial – review IT", "FFE699", 3),
    "compliant": ("Compliant", "C6EFCE", 2),
    "candidate": ("To verify (index)", "DDEBF7", 1),
    None: ("No related IT", "EDEDED", 0),
}
CHANGE = {"new": "New", "modified": "Changed", "unchanged": "Unchanged", None: ""}
CENTRAL = ("VA", "PI")  # central LEONI procedures (column PROC. WSD); WI / IC are the plant's own (PROC. WPTGU)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")


def _req_status(links):
    text = [l for l in links if l["source"] == "text" and l["verdict"] != "not_related"]
    if text:
        return max((l["verdict"] for l in text), key=lambda v: STATUS[v][2])
    return "candidate" if any(l["source"] == "catalog" for l in links) else None


def _table(ws, row, headers, widths):
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, WRAP, BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _row(ws, row, values, fill=None):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.alignment, c.border = WRAP, BORDER
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)


def _document_sections(store: Store, doc_id: int) -> OrderedDict:
    """Rows of the matrix in reading order: heading -> {page, top}, from the document outline down to the
    sub-section level (top-level headings are bold in the form), plus any passage section the outline lacks."""
    out = OrderedDict()
    for r in store.q("SELECT heading, level, page FROM outline WHERE doc_id=? AND level<=? ORDER BY idx",
                     (doc_id, MATRIX_LEVEL)):
        out.setdefault(r["heading"], {"page": r["page"], "top": r["level"] == 1})
    for r in store.q("SELECT section, page, top_section FROM chunks WHERE doc_id=? ORDER BY idx", (doc_id,)):
        if r["section"] and r["section"] not in out:
            out[r["section"]] = {"page": r["page"], "top": r["section"] == r["top_section"]}
    return out


def export_matrix(store: Store, doc_id: int, path=None) -> str:
    doc = store.one("SELECT * FROM documents WHERE id=?", (doc_id,))
    if not doc or doc["kind"] != "CSR":
        raise ValueError(f"document {doc_id} is not a customer document")
    catalog = {r["code"]: r for r in store.q("SELECT * FROM it_catalog")}
    it_docs = {d["code"]: d for d in store.q("SELECT code, title, meta_json FROM documents "
                                             "WHERE kind='IT' AND status='active'")}
    reqs = store.q("SELECT * FROM requirements WHERE doc_id=? ORDER BY page, id", (doc_id,))
    links = {}
    for l in store.q(f"SELECT * FROM req_it_links WHERE req_id IN (SELECT id FROM requirements WHERE doc_id=?)"
                     f" ORDER BY score DESC", (doc_id,)):
        links.setdefault(l["req_id"], []).append(l)

    # mock ITs (tools/mock_its.py): their content is invented, so every verdict based on them is marked
    synthetic = {c for c, d in it_docs.items() if json.loads(d["meta_json"] or "{}").get("synthetic")}

    def mark(code):
        return f"{code}*" if code in synthetic else code

    # "new" only means something when an earlier version of this document was ingested
    has_previous = bool(store.one("SELECT 1 FROM documents WHERE superseded_by=?", (doc_id,)))

    def base(code):
        return " ".join(code.split()[:2])

    def process_of(code):
        return catalog[base(code)]["process"] if base(code) in catalog else ""

    def it_label(code):
        """'WI 3157-11' -> 'IT 3157-11 Formação / Integração de colaboradores' (LEONI writes local WIs as IT)."""
        prefix, rest = code.split(" ", 1)
        title = it_docs[code]["title"] if code in it_docs else catalog.get(base(code), {"title": ""})["title"]
        return f"{'IT' if prefix == 'WI' else prefix} {rest}{'*' if code in synthetic else ''} {title}".strip()

    # ITs per requirement-section, ranked TF-IDF style: many requirements of the section point to the IT
    # (tf) and the IT is not one that shows up everywhere (idf). ITs with a conflict/gap come first.
    by_section = OrderedDict()
    for r in reqs:
        by_section.setdefault(r["section"] or "(no section)", []).append(r)
    counts, flagged = {}, {}
    for sec, rs in by_section.items():
        c = counts[sec] = Counter()
        flagged[sec] = {}
        for r in rs:
            for l in links.get(r["id"], []):
                if l["source"] == "text" and l["verdict"] == "not_related":
                    continue
                c[l["it_code"]] += 1
                if l["verdict"] in ("conflict", "partial"):
                    flagged[sec][l["it_code"]] = l["verdict"]
    df = Counter(code for c in counts.values() for code in c)
    n_sec = max(len(counts), 1)

    def ranked(sec):
        c = counts.get(sec, Counter())
        return sorted(c, key=lambda k: (k not in flagged[sec], -c[k] * math.log(1 + n_sec / df[k])))

    # section-level mapping (engine.map_sections), keyed by top-level heading
    section_its = {}
    for s_ in store.q("SELECT * FROM section_its WHERE doc_id=? ORDER BY id", (doc_id,)):
        section_its.setdefault(s_["section"], []).append(s_["it_code"])

    # section headings in the output language (translated once and cached; originals if the LLM is unavailable)
    try:
        label = section_labels(store, doc_id)
    except Exception:  # noqa: BLE001 - the export must work without the LLM
        label = {}

    def lbl(sec):
        return label.get(sec, sec)

    wb = load_workbook(TEMPLATE)

    # ---- Sheet 1: "Matrix" in LEONI's IT 3535-11 Anexo 5 format -------------------------------
    ws = wb["Matrix"]
    proto = {kind: [(copy.copy(ws.cell(r, c).font), copy.copy(ws.cell(r, c).border),
                     copy.copy(ws.cell(r, c).alignment), copy.copy(ws.cell(r, c).fill)) for c in range(1, 11)]
             for kind, r in (("top", FIRST_ROW), ("sub", FIRST_ROW + 1))}
    date = " / rev. ".join(dict.fromkeys(x for x in (doc["doc_date"], doc["version"]) if x)) or "-"
    ws["B4"] = f"Cliente:  {(doc['client'] or '').upper()}"
    ws["B5"] = f"Título:     {doc['title']}  ({doc['code']})"
    ws["B6"] = f"Publication Date:   {date}"
    note = (f"Generated by LEONI COMPLY-AI on {datetime.date.today():%d.%m.%Y}. PROC. WSD lists only the central "
            f"LEONI procedures loaded in the system: complete it manually.")
    if synthetic:
        note += " WARNING: includes SIMULATED ITs (*), generated for demonstration — not real LEONI ITs."
    ws["B7"] = note
    ws["B7"].font = Font(italic=True, size=9, color="C00000" if synthetic else "7F7F7F")

    row = FIRST_ROW
    for sec, info in _document_sections(store, doc_id).items():
        local = [c for c in ranked(sec) if c.split()[0] not in CENTRAL]
        if info["top"] and section_its.get(sec):
            its = section_its[sec] + [c for c in local if base(c) not in section_its[sec]]
            its = its[:max(config.MATRIX_ITS_PER_SECTION, len(section_its[sec]))]
        else:
            its = local[:3]
        its = list(dict.fromkeys(its))
        wsd = [c for c in ranked(sec) if c.split()[0] in CENTRAL][:3]
        values = {COL["section"]: lbl(sec), COL["page"]: info["page"],
                  COL["process"]: "\n".join(dict.fromkeys(p for p in map(process_of, its) if p)) or None,
                  COL["wsd"]: "\n".join(f"• {it_label(c)}" for c in wsd) or None,
                  COL["wptgu"]: "\n".join(f"• {it_label(c)}" for c in its) or None}
        for c, (font, border, align, fill) in enumerate(proto["top" if info["top"] else "sub"], start=1):
            cell = ws.cell(row, c, values.get(c))
            cell.font, cell.border, cell.alignment, cell.fill = font, border, align, fill
        row += 1
    ws.print_area = f"A1:J{max(row - 1, FIRST_ROW)}"

    # ---- Sheet 2: verdicts per requirement-section ---------------------------------------------
    ws = wb.create_sheet("Status by section")
    headers = ["Section", "Page", "Requirements", "ITs assessed / suggested", "Status",
               "Changes vs previous version"]
    _table(ws, 1, headers, [45, 8, 10, 60, 20, 25])
    for i, (sec, rs) in enumerate(by_section.items(), start=2):
        worst = max((_req_status(links.get(r["id"], [])) for r in rs), key=lambda v: STATUS[v][2])
        changes = ([CHANGE[r["change_type"]] for r in rs if r["change_type"] in ("new", "modified")]
                   if has_previous else [])
        its_txt = "\n".join(f"• {mark(c)} ({counts[sec][c]} req.)"
                            + (f" — {STATUS[flagged[sec][c]][0]}" if c in flagged[sec] else "")
                            for c in ranked(sec)[:config.MATRIX_ITS_PER_SECTION])
        ch_txt = ", ".join(f"{changes.count(x)} {x.lower()}" for x in dict.fromkeys(changes))
        _row(ws, i, [lbl(sec), rs[0]["page"], len(rs), its_txt, STATUS[worst][0], ch_txt], fill=STATUS[worst][1])

    # ---- Sheet 3: one line per requirement -------------------------------------------------
    ws = wb.create_sheet("Requirements")
    headers = ["ID", "Section", "Page", "Requirement", "Original text", "Category", "Obligation", "Parameters",
               "Change", "Change detail", "Status", "IT assessed", "Analysis", "IT evidence",
               "Recommended action", "Suggested ITs (index)", "LEONI process"]
    _table(ws, 1, headers, [6, 28, 7, 50, 50, 18, 12, 22, 12, 35, 18, 16, 45, 40, 40, 30, 14])
    for i, r in enumerate(reqs, start=2):
        ls = links.get(r["id"], [])
        st = _req_status(ls)
        text = [l for l in ls if l["source"] == "text" and l["verdict"] != "not_related"]
        text.sort(key=lambda l: -STATUS[l["verdict"]][2])
        cat = [l["it_code"] for l in ls if l["source"] == "catalog"]
        params = "; ".join(f"{p['name']}: {p['value']} {p['unit']}".strip()
                           for p in json.loads(r["parameters_json"] or "[]"))
        procs = {process_of(c) for c in [l["it_code"] for l in text] + cat} - {""}
        _row(ws, i, [r["id"], lbl(r["section"] or ""), r["page"], r["statement"], r["quote"], r["category"], r["obligation"],
                     params, CHANGE[r["change_type"]] if has_previous else "", r["change_note"] or "", STATUS[st][0],
                     "\n".join(mark(l["it_code"]) for l in text),
                     "\n".join(f"[{mark(l['it_code'])}] {l['explanation']}" for l in text),
                     "\n".join(l["evidence"] or "" for l in text),
                     "\n".join(l["action"] or "" for l in text if l["action"]),
                     "\n".join(f"{c} {catalog[c]['title']}" if c in catalog else c for c in cat),
                     ", ".join(sorted(procs))], fill=STATUS[st][1])

    # ---- Sheet 4: alerts -------------------------------------------------------------------
    ws = wb.create_sheet("Alerts")
    headers = ["ID", "Date", "Type", "Severity", "IT", "Title", "Detail", "Action", "Status"]
    _table(ws, 1, headers, [6, 18, 14, 10, 16, 45, 70, 45, 10])
    sev_fill = {"high": "F8CBAD", "medium": "FFE699", "low": "DDEBF7", "info": None}
    findings = store.q("""SELECT * FROM findings WHERE doc_id=? ORDER BY
                          CASE severity WHEN 'high' THEN 0 WHEN 'medium' THEN 1 WHEN 'low' THEN 2 ELSE 3 END, id""",
                       (doc_id,))
    for i, f in enumerate(findings, start=2):
        _row(ws, i, [f["id"], f["created_at"], f["type"], f["severity"], f["it_code"], f["title"], f["detail"],
                     f["action"], f["status"]], fill=sev_fill.get(f["severity"]))

    config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in f"{doc['client']}_{doc['code']}")
    path = path or config.EXPORT_DIR / f"Matriz_{safe}_v{doc['version'] or ''}.xlsx".replace(" ", "")
    wb.active = 0
    wb.save(path)
    return str(path)
